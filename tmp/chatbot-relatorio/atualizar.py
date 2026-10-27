from pathlib import Path
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
root=Path(r'C:/6Semestre/PLN/Atividade ChatBot I - Parte 2')
out=root/'Relatorio_Atividade_ChatBot_I_Parte_2_Iago_Yuri_Rossan.docx'
d=Document(); sec=d.sections[0]; sec.page_width=Cm(21);sec.page_height=Cm(29.7);sec.top_margin=Cm(1.8);sec.bottom_margin=Cm(1.8);sec.left_margin=Cm(2);sec.right_margin=Cm(2)
for name,size in [('Normal',11),('Title',24),('Subtitle',12),('Heading 1',17),('Heading 2',12)]:
 s=d.styles[name];s.font.name='Calibri';s.font.size=Pt(size);s.font.color.rgb=RGBColor(0,0,0);s.paragraph_format.space_after=Pt(7)
d.styles['Normal'].paragraph_format.line_spacing=1.08
for border in list(d.styles.element.iter(qn('w:pBdr'))):border.getparent().remove(border)
d.core_properties.author='Iago Yuri Rossan';d.core_properties.title='Atividade ChatBot I Parte 2';d.core_properties.subject='Implementação, resultados e roteiro de testes manuais'
f=sec.footer.paragraphs[0];f.alignment=2;f.add_run('Iago Yuri Rossan | ').font.size=Pt(9);field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');f._p.append(field)
def p(t):d.add_paragraph(t)
def h(t):d.add_heading(t,1)
def sub(t):d.add_heading(t,2)
def page(t):d.add_page_break();h(t)
def table(headers,rows,widths):
 t=d.add_table(rows=1,cols=len(headers));t.autofit=False
 for col,w in zip(t.columns,widths):col.width=Cm(w)
 for i,v in enumerate(headers):t.rows[0].cells[i].text=v
 for vals in rows:
  for c,v in zip(t.add_row().cells,vals):c.text=v
 for ri,row in enumerate(t.rows):
  pr=row._tr.get_or_add_trPr();pr.append(OxmlElement('w:cantSplit'))
  if ri==0:pr.append(OxmlElement('w:tblHeader'))
  for ci,c in enumerate(row.cells):
   c.width=Cm(widths[ci]);tc=c._tc.get_or_add_tcPr();b=OxmlElement('w:tcBorders')
   for edge in ['top','bottom','left','right']:
    e=OxmlElement('w:'+edge);e.set(qn('w:val'),'single');e.set(qn('w:sz'),'4');e.set(qn('w:color'),'D9D9D9');b.append(e)
   tc.append(b);margin=OxmlElement('w:tcMar')
   for edge in ['top','bottom','left','right']:
    e=OxmlElement('w:'+edge);e.set(qn('w:w'),'90');e.set(qn('w:type'),'dxa');margin.append(e)
   tc.append(margin)
   if ri==0:
    sh=OxmlElement('w:shd');sh.set(qn('w:fill'),'E8EDF3');tc.append(sh)
   for para in c.paragraphs:
    para.paragraph_format.space_after=Pt(2);para.paragraph_format.line_spacing=1.05
    for run in para.runs:run.font.size=Pt(10);run.bold=ri==0
 d.add_paragraph().paragraph_format.space_after=Pt(0)
def code(t):
 para=d.add_paragraph();para.paragraph_format.space_after=Pt(9)
 run=para.add_run(t);run.font.name='Consolas';run.font.size=Pt(9)
p('FATEC VOTORANTIM')
d.add_paragraph('Atividade ChatBot I Parte 2','Title')
d.add_paragraph('Implementação e roteiro de testes','Subtitle')
p('Iago Yuri Rossan\n6º semestre de DSM\nDesenvolvimento de Software Multiplataforma\nPLN — Processamento de Linguagem Natural\n26 de outubro de 2026')
h('1 Desenvolvimento da atividade')
p('Nesta atividade, ampliei dois chatbots a partir do repositório fornecido pela professora. Implementei novas regras no chatbot de suporte e novas intenções no chatbot baseado em PLN. Executei verificações das respostas e da interface e organizei um roteiro para repetir os testes manualmente.')
p('Utilizei o clone existente em C:\\chat, do repositório https://github.com/jylaila/chat. As alterações ficaram na pasta PLN/Atividade ChatBot I - Parte 2 do repositório 6Semestre. Mantive a estrutura simples da base: server.js, index.html, package.json e package-lock.json em cada projeto.')
sub('Chatbot baseado em regras')
p('No projeto RulesBased/chatbot-node, mantive login, planos e atendente e acrescentei horário, localização e promoções. A mensagem é convertida para minúsculas e verificada por if/else com includes. A primeira condição encontrada determina a resposta. O servidor utiliza a porta 3000.')
sub('Chatbot baseado em PLN')
p('No projeto PLNBased/chatbot-pln, mantive Natural e BayesClassifier. Acrescentei as intenções preco, ajuda e cancelamento, com três exemplos para cada uma. Junto dos oito exemplos originais, o treinamento contém 17 frases e seis intenções. A chamada classifier.train() treina o modelo ao iniciar o servidor, que utiliza a porta 3001.')
sub('Interface e organização final')
p('Melhorei o index.html dos dois projetos com fundo suave, painel centralizado, balões de mensagens e formulário adaptado para celular. O CSS permanece no próprio HTML. A estrutura final não contém scripts de testes nem pasta de evidências; as capturas estão incorporadas neste relatório. Os dados apresentados são fictícios e não há operações comerciais reais.')
page('2 Como executar e testar manualmente')
p('Para testar, utilizo Node.js com npm e mantenho dois terminais abertos, um para cada servidor. O ambiente usado nas verificações tinha Node.js 24.15.0, Express 5.1.0 e Natural 8.1.0.')
sub('Primeiro terminal para o chatbot de regras')
code('cd "C:\\6Semestre\\PLN\\Atividade ChatBot I - Parte 2\\RulesBased\\chatbot-node"\nnpm install\nnode server.js')
p('Após aparecer a mensagem de que o servidor está rodando, abro http://localhost:3000 no navegador. Digito uma pergunta no campo e clico em Enviar ou pressiono Enter. Comparo a resposta com a tabela da próxima página.')
sub('Segundo terminal para o chatbot de PLN')
code('cd "C:\\6Semestre\\PLN\\Atividade ChatBot I - Parte 2\\PLNBased\\chatbot-pln"\nnpm install\nnode server.js')
p('Abro http://localhost:3001. Envio as frases de teste e observo tanto o texto da resposta quanto a intenção mostrada entre parênteses, por exemplo, Bot (preco). Posso manter os dois chatbots abertos ao mesmo tempo, pois utilizam portas diferentes.')
sub('Quando instalar e reiniciar')
p('Executo npm install na primeira utilização ou quando as dependências mudarem. Nas próximas execuções, basta node server.js. Também posso iniciar com npm start. Mantenho o terminal aberto durante o teste e pressiono Ctrl+C para encerrar.')
p('Após mudar o server.js, interrompo o processo e inicio novamente. No PLN, isso também retreina o classificador com todos os addDocument. Após mudar apenas o HTML ou o CSS, atualizo o navegador com Ctrl+F5. Abro sempre a página pelo endereço localhost, e não diretamente como arquivo index.html.')
sub('Dificuldades comuns')
table(['Situação','Como resolver'],[['Cannot find module','Executar npm install na pasta do chatbot correspondente.'],['EADDRINUSE','A porta já está ocupada. Encerrar a execução anterior antes de iniciar outra.'],['Página não abre','Conferir se o terminal está aberto, se o servidor iniciou e se a porta é 3000 ou 3001.'],['Visual antigo','Atualizar com Ctrl+F5 para recarregar o HTML e o CSS.']],[4.4,12.6])
page('3 Testes do chatbot de regras')
p('Acesso http://localhost:3000 e envio uma frase por vez. Nos testes já executados, conferi as regras originais, as três regras novas, grafias sem acento, maiúsculas e uma mensagem sem palavra prevista. O resultado foi compatível com as respostas programadas.')
table(['Mensagem para testar','Resultado esperado'],[['Qual o horário?','Atendimento de segunda a sexta, das 8h às 18h.'],['Qual o horario?','A mesma resposta de horário, mesmo sem acento.'],['Qual a localização?','Nossa loja fica no Centro de Votorantim.'],['Qual o endereco?','A mesma resposta de localização.'],['Tem promoção?','10% de desconto para novos clientes.'],['Preciso de login','Orientação para clicar em Esqueci minha senha.'],['Quero um plano','Informação sobre os planos Básico e Premium.'],['Quero um atendente','Informação de que o atendimento é simulado.'],['Gosto de futebol','Desculpe, não entendi sua pergunta.']],[6.2,10.8])
p('A tabela funciona como roteiro de teste manual. Foram testadas também as variações “Qual o HORÁRIO?”, “Há promoções?”, “Tem promocao?” e “Bom tempo hoje”. A última recebeu a resposta padrão por não conter nenhuma palavra reconhecida.')
d.add_picture(r'C:/Users/Iago Rossan/AppData/Local/Temp/chat-layout-regras.png',width=Cm(10.8))
p('Figura 1 — Interface atual após o envio de “Qual o horário?”. O chatbot retornou o período de atendimento esperado.')
page('4 Testes do chatbot baseado em PLN')
p('Acesso http://localhost:3001 e comparo a intenção exibida com a tabela. As nove frases abaixo foram classificadas corretamente nas verificações HTTP realizadas durante o desenvolvimento. São casos simples do domínio usado no treinamento.')
table(['Mensagem testada','Intenção esperada'],[['oi','saudacao'],['acompanhar meu pedido','pedido_status'],['muito obrigado','agradecimento'],['qual o preço do plano','preco'],['quanto custa o plano','preco'],['preciso de ajuda agora','ajuda'],['quero suporte por favor','ajuda'],['quero cancelar minha assinatura','cancelamento'],['como solicitar cancelamento','cancelamento']],[11,6])
p('Para preco, espero os valores de R$ 29,90 para o Básico e R$ 59,90 para o Premium. Para ajuda, espero a solicitação de descrição da dúvida. Para cancelamento, espero a orientação de entrar em contato com o atendimento. Nenhuma dessas respostas realiza uma operação real.')
p('Para conferir o retreinamento, posso adicionar uma frase a uma intenção com classifier.addDocument, antes de classifier.train(), e reiniciar o servidor. Depois envio uma frase relacionada e observo a intenção retornada. Os testes deste relatório foram feitos após incluir as três novas intenções.')
d.add_picture(r'C:/Users/Iago Rossan/AppData/Local/Temp/chat-layout-pln.png',width=Cm(10.8))
p('Figura 2 — Interface atual após o envio de “quero cancelar minha assinatura”. A intenção identificada foi cancelamento, conforme esperado.')
page('5 Resultados observados e conclusão')
p('Durante o desenvolvimento, executei verificações HTTP e de navegador. Após a simplificação, removi os arquivos auxiliares de teste para manter o projeto próximo da estrutura original. Os resultados abaixo registram o que foi executado; o roteiro das páginas anteriores permite repetir a conferência manualmente.')
table(['Verificação realizada','Resultado observado'],[['Regras por HTTP','12 verificações aprovadas: 10 mensagens, página inicial e rejeição de mensagem vazia.'],['PLN por HTTP','11 verificações aprovadas: 9 mensagens, página inicial e rejeição de mensagem vazia.'],['Interface simples no navegador','6 verificações aprovadas: três novas regras e três novas intenções.'],['Interface atual após alterar o visual','Envio e resposta de horário no chatbot de regras; envio e resposta de cancelamento no PLN. Ambos funcionaram.'],['Layout de celular','As duas interfaces foram conferidas em largura de 390 pixels, sem transbordamento horizontal.']],[6,11])
p('A execução HTTP totalizou 23 verificações aprovadas. O executor do Node registrou 25 testes ao incluir os dois testes principais que agrupavam essas verificações. As verificações de navegador foram realizadas com Playwright e Microsoft Edge. As capturas deste documento correspondem ao visual atual.')
sub('Conferências adicionais durante o teste manual')
p('Ao clicar em Enviar com o campo vazio, espero que nenhuma nova mensagem seja adicionada. Ao enviar uma mensagem válida, espero vê-la no balão do usuário e, em seguida, receber a resposta do bot. No PLN, também confiro a intenção entre parênteses. Se interromper o servidor, uma nova tentativa de envio deverá apresentar uma mensagem de erro de comunicação.')
sub('Diferença entre as abordagens')
p('No chatbot de regras, a resposta depende das palavras previstas no if/else. Uma pergunta com outra forma de escrever pode não ser reconhecida. Se uma frase contiver palavras de mais de uma regra, vence a primeira condição na ordem do código; por exemplo, uma frase com plano e promoção aciona a regra de plano.')
p('No PLN, o classificador utiliza os exemplos para escolher uma das intenções aprendidas. Ele pode reconhecer variações, mas também pode classificar incorretamente uma frase desconhecida. O Bayes sempre escolhe uma classe: a resposta padrão do switch não representa uma garantia de rejeição de assuntos fora do domínio.')
sub('Conclusão')
p('Concluí a inclusão das regras e intenções solicitadas, retreinei o classificador e conferi as respostas esperadas. Os testes demonstraram o funcionamento dos casos escolhidos. O projeto permanece simples e adequado à atividade, sem memória de conversa ou integração com sistemas reais. O acerto nesses exemplos não garante acerto em qualquer mensagem.')
p('Repositório-base: https://github.com/jylaila/chat\nRepositório da disciplina: https://github.com/IagoYuriRossan/6Semestre')
d.save(out)

